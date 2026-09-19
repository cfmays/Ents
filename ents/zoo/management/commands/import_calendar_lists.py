"""
Preload one calendar (ASG) per tab of "Enrichment Calendars v4.xlsx"
(or, with --aquarium, "Aquarium Enrichment Calendars V4.xlsx").

For each calendar tab, reads its "<tab> List" sheet:
  - column A/E rows 1-12  -> special concerns
  - column I              -> Animal choices (e.g. "Rocky/ Raza/ Nety")
  - column J              -> behavior goals
  - column A/B/C, below the "APPROVED NON-FOOD ENRICHMENT" heading (and its column headers) -> approved NON-FOOD items, comments, rate
  - column E/F/G, same rows                                         -> approved FOOD items, comments, rate
and the "Notes:" cell of the calendar tab itself.

Items are matched to the master list by exact name (extra spaces and upper/lower case ignored).
Anything that doesn't match is NOT applied; it is written to the report file.

Run a dry run first (rolls everything back):
  python3 manage.py import_calendar_lists "../Enrichment Calendars v4.xlsx" --dry-run
"""
import re
from collections import defaultdict

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from ents.models import Enrichment
from zoo.models import ASG, ASGApprovedItem, Animal, BehaviorGoal, SpecialConcern, String

STRING_TABS = {
    'Komodo String': [
        'Galapagos tortoise', 'Aldabra tortoise', 'Aruba Island Rattlesnake', 'Large Venomous Snakes', 'Komodo',
    ],
    'Ambassador Reptiles/ Amphibian/ Inverts': [
        'Corn Snake', 'Dart Frog', 'Box Turtle', 'Uromastyx', 'Skink Blue Tongue', 'Bearded Dragon',
        'Skink Prenhensile Tailed', 'Burmese Star Torts', 'Legless Lizard', 'Toad', 'Whites tree frog',
        'Black rat snake', 'Tarantulas',
    ],
    'Birdhouse String': [
        'Emerald Tree Boa', 'Rosy Boa', 'Ball Python', 'Millipedes, Roaches', 'Barn Owl', 'Kook', 'Macaw',
        'Trumpeter Hornbill',
    ],
    'Ambassador Mammal/ Bird': [
        'SH Armadillo - Malcolm', 'SH Amadillo - Alice', 'SH Armadillo- Rizzo', 'Ferret', 'Virginia Opossum',
        'Prehensile Tailed Porcupine', 'Chinchilla', 'Wood Rat',
    ],
    'Lemur String': [
        'RR Lemur', 'RT Lemur', 'Radiated Tortoise', 'Fennec Fox Female', 'Fennec Fox', 'Sand Cat', 'Fishing Cat',
        'Serval Tut', 'Serval Kira', 'Serval Juvenile',
    ],
    'Fossa (was Gibbon) String': ['Binturong', 'Howler Monkey', 'Gibbon', 'Fossa'],
    'Tiger/ Wolf String': ['Tiger', 'Maned Wolf', 'Red Panda'],
    'Barn Ambassador': ['HorseBurro', 'Goat', 'Chickens'],
    'Hippo / Okapi String': [
        'Okapi', 'Southern Ground Hornbill', 'Flamingo', 'Hippo', 'Hippo Fish', 'Cassowary', 'Bluegill',
    ],
}
AQUARIUM_STRING_TABS = {
    'Aquarium Mammal': [
        'Tamarin - Titi - Caiman', 'Sloth', 'AB1- pool', 'Tamarin- Quarantine', 'Otters', 'Otter- Quarantine',
    ],
    'Aquarium Bird/ Reptile': [
        'Penguins', 'Penguins- Quarantine', 'Penguin Pool', 'Diamondback Terrapin', 'Alligator Snapping Turtle',
    ],
    'Aquarium Fish': [
        'Rogue', 'Surfer', 'Quarantine Tanks', 'AQ Universal', 'All AQ Enrichment', 'AQ Holiday Enrichment',
    ],
}
AQUARIUM_TANK_STRING = 'Aquarium Fish'  # every "Tank ..." tab
# Obsolete per the user; not imported.
SKIPPED_TABS = {'Tamarin - Titi - Caiman - Sloth'}


def string_for_tab(name, aquarium):
    if aquarium and name.startswith('Tank '):
        return AQUARIUM_TANK_STRING
    tabs = {tab: string for string, tabs in (AQUARIUM_STRING_TABS if aquarium else STRING_TABS).items() for tab in tabs}
    return tabs.get(name)


# Calendar tab -> list sheet, where the sheet isn't named "<tab> List".
LIST_SHEET_OVERRIDES = {'Serval Kira': 'Serval List'}

# Sheet spellings the user confirmed are the same item as a master-list name.
ITEM_ALIASES = {
    'moving exhibit fortuner': 'Moving exhibit furniture',
    'ball-10” - hard plastic': 'Ball- 10”- Hard Plastic',
    'ball- large yellow bouy- 2in hole': 'Ball- Large Yellow Buoy- 2in hole',
    'slow feeder- no rubber base': 'Slow Feeder- Any With No Rubber Base',
}

# Typed over the "Special Concerns" header cell in a few list sheets.
JUNK_CONCERNS = {'porg', 'sec'}
HEADER_TEXT = {'special concerns', 'more special concerns', 'animals', 'behavior goals'}


def clean(value):
    """Collapse whitespace; '' for empty cells."""
    return re.sub(r'\s+', ' ', str(value)).strip() if value is not None else ''


def clean_rate(value):
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean(value)


def calendar_name(title):
    return clean(title.replace('"', ''))


class Command(BaseCommand):
    help = 'Preload calendars (ASGs) from Enrichment Calendars v4.xlsx'

    def add_arguments(self, parser):
        parser.add_argument('xlsx_path')
        parser.add_argument('--aquarium', action='store_true', help='Use the aquarium tab -> String map.')
        parser.add_argument('--dry-run', action='store_true', help='Do everything, then roll back.')
        parser.add_argument('--report', default='calendar_import_report.txt', help='Where to write the report.')

    def handle(self, *args, **options):
        aquarium = options['aquarium']
        wb = openpyxl.load_workbook(options['xlsx_path'], data_only=True)
        titles = wb.sheetnames
        first_list = next(i for i, t in enumerate(titles) if t.endswith((' List', ' Lis')))
        calendar_titles = titles[:first_list]
        list_sheets = {clean(t): t for t in titles if t.endswith((' List', ' Lis'))}

        # Every calendar must have a String and a list sheet; stop before writing anything if not.
        plan = []
        for title in calendar_titles:
            name = calendar_name(title)
            if name in SKIPPED_TABS:
                continue
            if string_for_tab(name, aquarium) is None:
                raise CommandError(f'No String chosen for calendar tab "{name}".')
            if name in LIST_SHEET_OVERRIDES:
                matches = [LIST_SHEET_OVERRIDES[name]]
            else:
                list_key = clean((name + ' List').replace('"', ''))
                matches = [t for key, t in list_sheets.items()
                           if clean(re.sub(r' Lis(t)?$', '', key).replace('"', '')) + ' List' == list_key]
            if len(matches) != 1 or matches[0] not in titles:
                raise CommandError(f'Expected one list sheet for "{name}", found {matches}.')
            plan.append((title, name, matches[0]))

        strings = {s.name: s for s in String.objects.all()}
        needed = {string_for_tab(name, aquarium) for _, name, _ in plan}
        missing = needed - set(strings)
        if missing:
            raise CommandError(f'Strings not found: {sorted(missing)}')

        items_by_name = defaultdict(list)
        for item in Enrichment.objects.select_related('category'):
            items_by_name[clean(item.name).lower()].append(item)
        goals_by_lower = {clean(g.name).lower(): g for g in BehaviorGoal.objects.all()}

        report = []
        totals = defaultdict(int)
        with transaction.atomic():
            for title, name, list_title in plan:
                cal_ws, list_ws = wb[title], wb[list_title]
                asg, created = ASG.objects.get_or_create(
                    name=name, defaults={'string': strings[string_for_tab(name, aquarium)], 'notes': self.read_notes(cal_ws)},
                )
                counts = defaultdict(int)

                for text in self.read_special_concerns(list_ws, name, report):
                    concern, _ = SpecialConcern.objects.get_or_create(text=text)
                    asg.special_concerns.add(concern)
                    counts['concerns'] += 1

                for row in range(2, 13):
                    text = clean(list_ws[f'I{row}'].value)
                    if text and text.lower() not in HEADER_TEXT:
                        Animal.objects.get_or_create(asg=asg, name=text)
                        counts['animals'] += 1

                for row in range(2, list_ws.max_row + 1):
                    text = clean(list_ws[f'J{row}'].value)
                    if not text or text.lower() in HEADER_TEXT:
                        continue
                    goal = goals_by_lower.get(text.lower())
                    if goal is None:
                        goal = BehaviorGoal.objects.create(name=text)
                        goals_by_lower[text.lower()] = goal
                        report.append(f'NEW GOAL "{text}" (first seen on {name})')
                    asg.behavior_goals.add(goal)
                    counts['goals'] += 1

                for cols, is_food in ((('A', 'B', 'C'), False), (('E', 'F', 'G'), True)):
                    self.read_items(asg, list_ws, cols, is_food, items_by_name, counts, report)

                for key, value in counts.items():
                    totals[key] += value
                self.stdout.write(
                    f'{name:32} {asg.string.name:40} {"new" if created else "existing":8} '
                    f'animals={counts["animals"]} concerns={counts["concerns"]} goals={counts["goals"]} '
                    f'nonfood={counts["nonfood"]} food={counts["food"]} UNMATCHED={counts["unmatched"]}'
                )

            if options['dry_run']:
                transaction.set_rollback(True)

        with open(options['report'], 'w') as f:
            f.write('\n'.join(report) + '\n')
        self.stdout.write(self.style.SUCCESS(
            f'{"DRY RUN (rolled back): " if options["dry_run"] else ""}{len(plan)} calendars, '
            f'{totals["nonfood"]} non-food + {totals["food"]} food items applied, '
            f'{totals["unmatched"]} rows NOT applied. Report: {options["report"]}'
        ))

    def first_item_row(self, list_ws):
        """First item row: below the 'APPROVED NON-FOOD ENRICHMENT' heading, skipping the column-header row if any."""
        for row in range(10, 25):
            if clean(list_ws[f'A{row}'].value).upper() == 'APPROVED NON-FOOD ENRICHMENT':
                has_header = clean(list_ws[f'A{row + 1}'].value).lower().endswith('category fulfillment')
                return row + 2 if has_header else row + 1
        raise CommandError(f'No "APPROVED NON-FOOD ENRICHMENT" heading found in {list_ws.title}')

    def read_notes(self, cal_ws):
        for row in range(1, 7):
            if clean(cal_ws[f'A{row}'].value).lower() == 'notes:':
                for cell in cal_ws[row][1:]:
                    if clean(cell.value):
                        return str(cell.value).strip()
        return ''

    def read_special_concerns(self, list_ws, name, report):
        for col in ('A', 'E'):
            for row in range(1, 13):
                text = clean(list_ws[f'{col}{row}'].value)
                if not text or text.lower() in HEADER_TEXT:
                    continue
                if text.lower() in JUNK_CONCERNS:
                    report.append(f'SKIPPED stray text "{text}" in {col}{row} of {name} list')
                    continue
                yield text

    def read_items(self, asg, list_ws, cols, is_food, items_by_name, counts, report):
        name_col, comment_col, rate_col = cols
        kind = 'food' if is_food else 'nonfood'
        merged = {}
        for row in range(self.first_item_row(list_ws), list_ws.max_row + 1):
            item_name = clean(list_ws[f'{name_col}{row}'].value)
            if not item_name:
                continue
            key_name = clean(ITEM_ALIASES.get(item_name.lower(), item_name)).lower()
            found = items_by_name.get(key_name, [])
            if len(found) != 1:
                counts['unmatched'] += 1
                why = 'no exact match in master list' if not found else 'several master-list items share this name'
                report.append(f'NOT APPLIED [{asg.name} {name_col}{row}] "{item_name}": {why}')
                continue
            comment = clean(list_ws[f'{comment_col}{row}'].value)
            rate = clean_rate(list_ws[f'{rate_col}{row}'].value)
            key = found[0].id
            if key in merged:
                old_comment, old_rate = merged[key]
                if comment and comment != old_comment:
                    comment = f'{old_comment} | {comment}' if old_comment else comment
                    report.append(f'MERGED differing comments for duplicate "{item_name}" in {asg.name}')
                merged[key] = (comment or old_comment, rate or old_rate)
            else:
                merged[key] = (comment, rate)

        for item_id, (comment, rate) in merged.items():
            ASGApprovedItem.objects.update_or_create(
                asg=asg, item_id=item_id, is_food=is_food,
                defaults={'comments': comment[:500], 'rate': rate[:100]},
            )
            counts[kind] += 1
